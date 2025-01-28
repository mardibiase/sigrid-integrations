from pptx.text.text import _Paragraph
from pptx.oxml.text import CT_TextParagraph
from src.report_generator.slide_utils import SlideUtils

class TestSlideUtils:

    def test_merge_similar_runs(self):
        p = _Paragraph(CT_TextParagraph(), None)
        r1 = p.add_run()
        r1.text = "aap"
        f1 = r1.font
        f1.bold = True

        r2 = p.add_run()
        r2.text = "noot"
        f2 = r2.font
        f2.bold = True
        
        SlideUtils.merge_runs_with_same_formatting(p)
        
        assert len(p.runs) == 1
        assert p.text == "aapnoot"

    def test_do_not_merge_different_runs(self):
        p = _Paragraph(CT_TextParagraph(), None)
        r1 = p.add_run()
        r2 = p.add_run()
        r1.text = "aap"
        r2.text = "noot"
        f1 = r1.font
        f1.bold = True
        f2 = r2.font
        f2.bold = False
        
        SlideUtils.merge_runs_with_same_formatting(p)

        assert len(p.runs) == 2
        assert p.runs[0].text == "aap"
        assert p.runs[1].text == "noot"